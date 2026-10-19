"""Builds REPORT.docx and REPORT.pdf from REPORT.md plus the latest results.

Usage: python evaluate.py && python docs/build_report.py   (needs pandoc; LibreOffice for the PDF)
"""

import csv
import subprocess
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

rows = list(csv.DictReader(open(ROOT / "results" / "test_results.csv")))
table = ["", "| ID | Group | Input | Expected | Actual | Result |", "|---|---|---|---|---|---|"]
for r in rows:
    text = r["input"] if len(r["input"]) <= 70 else r["input"][:67] + "..."
    text = text.replace("|", "/").replace("<", "\\<")  # keep pandoc from reading "</system>" as HTML
    table.append(f"| {r['id']} | {r['group']} | {text} | {r['expected']} | {r['actual']} | {r['status']} |")

full = (ROOT / "REPORT.md").read_text() + "\n".join(table) + "\n"
build = ROOT / "docs" / "_report_build.md"
build.write_text(full)
docx_path = ROOT / "REPORT.docx"
subprocess.run(["pandoc", str(build), "-o", str(docx_path), "--resource-path", str(ROOT)], check=True)
build.unlink()

# Pandoc gives every table equal, narrow columns. Size columns by content and use a smaller font.
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

doc = Document(docx_path)
usable = Cm(16.5)
for table in doc.tables:
    ncols = len(table.columns)
    lengths = [max(14, max(min(len(row.cells[i].text), 60) for row in table.rows)) for i in range(ncols)]
    total = sum(lengths)
    widths = [int(usable * n / total) for n in lengths]
    tbl_pr = table._tbl.tblPr
    for el in tbl_pr.findall(qn("w:tblW")):
        el.set(qn("w:type"), "dxa")
        el.set(qn("w:w"), str(int(usable.twips if hasattr(usable, "twips") else usable / 635)))
    for grid_col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(int(w / 635)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
doc.save(docx_path)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(ROOT), str(ROOT / "REPORT.docx")],
               check=True, capture_output=True)
print("wrote REPORT.docx and REPORT.pdf")
